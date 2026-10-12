#!/usr/bin/env python3
"""Genera borradores de specs SDD (una por HU y por tarea técnica) desde el backlog del proyecto.

Fuentes:
  --docx  Product_Backlog_Knock.docx      historia completa y criterios de aceptación
  --xlsx  Backlog_Planeacion_Knock.xlsx   épica, módulo, responsable, sprint, capas, plan de pruebas,
                                          trazabilidad y tareas técnicas

Cada criterio se vuelve un caso TC-<ID>-KK (mismo número que el plan de pruebas del proyecto).
Solo se crean las specs que tocan este repositorio (según las capas) y nunca se sobrescribe una spec
existente salvo con --force (y nunca una Aprobada o Firmada).

Requiere: pip install python-docx openpyxl
Uso: python scripts/import_backlog.py --docx ruta/Product_Backlog_Knock.docx --xlsx ruta/Backlog_Planeacion_Knock.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CONFIG = json.loads((ROOT / ".agents" / "knock.json").read_text(encoding="utf-8"))
REPO = CONFIG["repo"]
OTHER = "knock-frontend" if REPO == "knock-backend" else "knock-backend"
TEST_TYPES = {"UB": "Unitarias backend (xUnit)", "UM": "Unitarias app (flutter_test)", "W": "Widget (flutter_test)",
              "INT": "Integración (WebApplicationFactory + Testcontainers)", "ACC": "Aceptación (un caso TC por criterio)",
              "NEG": "Seguridad negativa (acceso ajeno, roles, límites)", "CON": "Contrato (OpenAPI)", "MUT": "Mutación",
              "E2E": "Extremo a extremo (integration_test / Playwright)", "PERF": "Rendimiento (k6)", "USA": "Usabilidad (SUS)",
              "A11Y": "Accesibilidad (contraste, etiquetas, texto escalable)"}
MODULE_SKILLS = {"identity": ["auth-firebase-jwt"], "catalog": ["postgis-geo"], "requests": ["events-outbox"],
                 "messaging": ["realtime-signalr"], "reputation": ["ai-features"], "trust": ["security-core"],
                 "notifications": ["events-outbox"], "assistant": ["ai-features"]}
BACK_HINT = re.compile(r"\b(40[0-9]|409|422|429|500|servidor|responde|se guarda|queda|auditor|evento|token|hash|base de datos|"
                       r"instancias|proceso|se crea|se borra|se anonimiza|registro|cifrad|endpoint|api|se envía|llega una sola)\b", re.I)
FRONT_HINT = re.compile(r"\b(veo|toco|pantalla|aparece|se muestra|botón|mensaje «|abro|deslizo|escribo|elijo|subo|marco)\b", re.I)


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")[:48].rstrip("-")


def read_docx(path: Path) -> dict[str, dict]:
    import docx
    stories = {}
    for table in docx.Document(path).tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            m = re.search(r"\b(HU-\d{3})\b", " ".join(cells[:2]))
            if not m:
                continue
            text = cells[2]
            title, _, story = text.partition(":")
            crit = [re.sub(r"^[-•]\s*", "", l).strip() for c in cells[4:] for l in c.split("\n") if re.match(r"^[-•]?\s*Dado", l.strip())]
            stories[m.group(1)] = {"title": title.strip(), "story": story.strip(), "criteria": crit}
    return stories


def read_xlsx(path: Path) -> tuple[dict, dict, dict, list]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def table(sheet: str, header_key: str) -> list[dict]:
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        start = next(i for i, r in enumerate(rows) if r and header_key in [str(c) for c in r if c is not None])
        head = [str(c).strip() if c is not None else "" for c in rows[start]]
        return [dict(zip(head, r)) for r in rows[start + 1:] if r and any(c is not None for c in r)]

    backlog = {r["ID"]: r for r in table("Backlog", "Historia de usuario") if str(r.get("ID", "")).startswith("HU-")}
    plan = {r["ID"]: r for r in table("Plan de pruebas", "Casos de prueba") if str(r.get("ID", "")).startswith("HU-")}
    trace = {}
    for r in table("Trazabilidad", "Casos de prueba"):
        m = re.match(r"(HU-\d{3})", str(r.get("Historia", "")))
        if m:
            trace[m.group(1)] = str(r.get("Fuente (DECISIONES / FUNCIONES / SEGURIDAD)", "") or "")
    tasks = [r for r in table("Tareas técnicas", "Tarea") if str(r.get("ID", "")).startswith("TT-")]
    return backlog, plan, trace, tasks


def repo_for_layers(layers: str) -> set[str]:
    repos = set()
    if re.search(r"\b(Back|BD)\b", layers):
        repos.add("knock-backend")
    if re.search(r"\b(Mobile|Panel)\b", layers):
        repos.add("knock-frontend")
    return repos


def verification(criterion: str, layers: str) -> str:
    back, front = bool(BACK_HINT.search(criterion)), bool(FRONT_HINT.search(criterion))
    if REPO == "knock-backend":
        if back or not front:
            return "Integración / seguridad negativa (por definir)"
        return f"(otro repo) Widget/E2E en {OTHER}" if re.search(r"Mobile|Panel", layers) else "Integración (por definir)"
    if front or not back:
        return "Widget / E2E (por definir)"
    return f"(otro repo) Integración en {OTHER} · E2E aquí si muestra el resultado"


def write(folder: Path, files: dict[str, str], force: bool) -> bool:
    spec = folder / "spec.md"
    if spec.exists():
        state = re.search(r"\*\*Estado:\*\*\s*(\w+)", spec.read_text(encoding="utf-8"))
        if not force or (state and state.group(1) in ("Aprobada", "Firmada")):
            return False
    for name, content in files.items():
        (folder / name).parent.mkdir(parents=True, exist_ok=True)
        (folder / name).write_text(content, encoding="utf-8")
    return True


def story_files(hid: str, s: dict, b: dict, plan: dict, sources: str) -> tuple[str, dict[str, str]]:
    layers = str(b.get("Capas", ""))
    module = str(b.get("Módulo", "")).split("/")[0].strip()
    slug = (re.search(r"KNOCK-<N>-(.+)$", str(b.get("Rama Git", ""))) or [None, slugify(s["title"])])[1]
    repos = repo_for_layers(layers)
    other = ", ".join(sorted(repos - {REPO})) or "—"
    skills = ["story-workflow", "reuse-before-create"] + MODULE_SKILLS.get(module.lower(), []) + CONFIG.get("default_skills", [])
    rows = "\n".join(f"| TC-{hid}-{i:02d} | {c.replace('|', '/')} | {verification(c, layers)} |" for i, c in enumerate(s["criteria"], 1))
    relevant = CONFIG["test_types"]
    plan_rows = "\n".join(
        f"| {TEST_TYPES[k]} | {plan.get(k, '—') if k in plan else '—'} | {'' if plan.get(k) != 'No' else 'Según el plan de pruebas del proyecto'} |"
        for k in relevant if k in TEST_TYPES)
    sprint = str(b.get("Sprint", "")).replace("Sprint ", "S")
    spec = f"""# {hid} — {s['title']}

> **Jira:** KNOCK-<N> · **Épica:** {b.get('Épica', '')} · **Módulo:** {b.get('Módulo', '')} · **Responsable:** {b.get('Responsable', '')} · **Sprint:** {sprint} · **Puntos:** {b.get('Story points', '')} · **Capas:** {layers}
> **Repositorio:** {REPO} · **Otra parte en:** {other}
> **Fuentes:** {sources or 'Product_Backlog_Knock.docx'} · Ficha: {b.get('Documento', '')}
> **Skills:** {', '.join(dict.fromkeys(skills))}
> **Estado:** Borrador

## 1. Objetivo
{s['story']}

Parte de este repositorio: {CONFIG['layer_scope']}. Nota del backlog: {b.get('Comentarios de seguimiento', '') or '—'}

## 2. Historia y casos de prueba
**{hid}** — {s['story']}

| Caso | Criterio (Dado / Cuando / Entonces) | Verificación |
|---|---|---|
{rows}

## 3. Alcance / No-alcance
- **In (este repositorio):** por refinar con `spec-writer` a partir de la ficha {b.get('Documento', '')}.
- **Out:** lo que corresponde a {other}; funciones de prioridad Media no incluidas en esta historia.

## 4. Contratos API y eventos
Por refinar: endpoints en `docs/endpoints.md`, eventos en `docs/events.md`{', contrato vendorizado en contracts/' if REPO == 'knock-frontend' else ''}.

## 5. Modelo de datos
Por refinar.

## 6. Seguridad, privacidad y amenazas
| Amenaza | Control | Caso que lo prueba |
|---|---|---|
| Por refinar (STRIDE proporcional al riesgo) | | |

## 7. Pantallas
{'Por refinar: pantalla(s) de la maqueta, estados cargando/vacío/error/datos, textos i18n y accesibilidad.' if REPO == 'knock-frontend' else '—'}

## 8. Pruebas, gates y límites
**Plan de pruebas** (del archivo Backlog y planeación; meta de mutación: {plan.get('Meta mutación', '—')}):

| Tipo | ¿Aplica? | Motivo / herramienta |
|---|---|---|
{plan_rows}

**Gates:** ver `AGENTS.md §5`.
**Límites y dudas abiertas:** refinar alcance por capa y amenazas antes de aprobar.

## Aprobación
- [ ] Revisor 1 — nombre: ________ fecha: ________
- [ ] Sin conflicto con la constitución
- **Estados:** Borrador (agente) → Aprobada (persona; habilita el código) → Firmada (tras el merge)
"""
    layer_tasks = [l.strip().replace(" web", "") for l in layers.split("+") if l.strip().replace(" web", "") in CONFIG["layers"]]
    tasks = "\n\n".join(f"## T{i} — {hid}-[{l}]\n- **Crear/Modificar:** por definir en plan.md\n- **Verificar:** casos de §2 de esta capa\n- **Agente:** {CONFIG['layers'].get(l, 'el constructor de la capa + traceability-clerk')}"
                        for i, l in enumerate(layer_tasks + ["Pruebas"], 1) if l == "Pruebas" or l in CONFIG["layers"])
    files = {
        "spec.md": spec,
        "plan.md": (ROOT / ".specify/templates/plan-template.md").read_text(encoding="utf-8").replace("<ID>", hid).replace("<Título>", s["title"]),
        "task.md": f"# {hid} — tareas ({REPO})\n\n{tasks}\n",
        "checklists/gates.md": (ROOT / ".specify/templates/checklist-template.md").read_text(encoding="utf-8").replace("<ID>", hid),
    }
    return f"{hid}-{slug}", files


def task_files(t: dict) -> tuple[str, dict[str, str]]:
    tid = t["ID"]
    slug = slugify(str(t["Tarea"]))[:40].rstrip("-")
    spec = f"""# {tid} — {t['Tarea']}

> **Jira:** KNOCK-<N> · **Área:** {t.get('Área', '')} · **Responsable:** {t.get('Responsable', '')} · **Sprint:** {str(t.get('Sprint', '')).replace('Sprint ', 'S')} · **Puntos:** {t.get('Pts', '')}
> **Repositorio:** {REPO} · **Ficha:** {t.get('Documento', '')}
> **Skills:** story-workflow, {', '.join(CONFIG.get('default_skills', []))}
> **Estado:** Borrador

## 1. Objetivo
{t['Tarea']}.

## 2. Historia y casos de prueba
**{tid}** — Tarea técnica (no es historia de usuario).

| Caso | Criterio (Dado / Cuando / Entonces) | Verificación |
|---|---|---|
| TC-{tid}-01 | Terminada cuando: {t.get('Terminada cuando', '')} | Por definir |

## 3. Alcance / No-alcance
- **In (este repositorio):** por refinar con la ficha {t.get('Documento', '')}.
- **Out:** —

## 4. Contratos API y eventos
—

## 5. Modelo de datos
—

## 6. Seguridad, privacidad y amenazas
| Amenaza | Control | Caso que lo prueba |
|---|---|---|
| Por refinar | | |

## 7. Pantallas
—

## 8. Pruebas, gates y límites
**Plan de pruebas:** por definir según la tarea.
**Gates:** ver `AGENTS.md §5`.
**Límites y dudas abiertas:** —

## Aprobación
- [ ] Revisor 1 — nombre: ________ fecha: ________
- [ ] Sin conflicto con la constitución
- **Estados:** Borrador (agente) → Aprobada (persona; habilita el código) → Firmada (tras el merge)
"""
    files = {
        "spec.md": spec,
        "plan.md": (ROOT / ".specify/templates/plan-template.md").read_text(encoding="utf-8").replace("<ID>", tid).replace("<Título>", str(t["Tarea"])),
        "task.md": f"# {tid} — tareas\n\n## T1 — {t['Tarea']}\n- **Verificar:** {t.get('Terminada cuando', '')}\n",
        "checklists/gates.md": (ROOT / ".specify/templates/checklist-template.md").read_text(encoding="utf-8").replace("<ID>", tid),
    }
    return f"{tid}-{slug}", files


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--force", action="store_true", help="regenera borradores existentes (nunca Aprobadas ni Firmadas)")
    args = parser.parse_args()
    stories = read_docx(args.docx)
    backlog, plan, trace, tasks = read_xlsx(args.xlsx)
    out = ROOT / "specs"
    index, created = [], 0
    for hid in sorted(backlog, key=lambda x: int(x[3:])):
        b = backlog[hid]
        if REPO not in repo_for_layers(str(b.get("Capas", ""))) or hid not in stories:
            continue
        name, files = story_files(hid, stories[hid], b, plan.get(hid, {}), trace.get(hid, ""))
        created += write(out / name, files, args.force)
        index.append((str(b.get("Sprint", "")), hid, stories[hid]["title"], str(b.get("Módulo", "")), str(b.get("Responsable", "")), str(b.get("Story points", "")), str(b.get("Prioridad", "")), name))
    for t in tasks:
        if t["ID"] not in CONFIG["technical_tasks"]:
            continue
        name, files = task_files(t)
        created += write(out / name, files, args.force)
        index.append((str(t.get("Sprint", "")), t["ID"], str(t["Tarea"]), str(t.get("Área", "")), str(t.get("Responsable", "")), str(t.get("Pts", "")), "Técnica", name))
    index.sort(key=lambda r: (int(re.sub(r"\D", "", r[0]) or 99), r[1]))
    lines = [f"# Specs — {REPO}", "", "> Generado por `scripts/import_backlog.py`. Cada carpeta: spec.md + plan.md + task.md + checklists/gates.md.",
             "> Estado de cada spec: dentro de su `spec.md`. Flujo: Borrador → Aprobada → Firmada.", "",
             "| Sprint | ID | Título | Módulo / área | Responsable | Pts | Prioridad | Carpeta |", "|---|---|---|---|---|---|---|---|"]
    lines += [f"| {r[0].replace('Sprint ', 'S')} | {r[1]} | {r[2]} | {r[3]} | {r[4].split(' ')[0]} | {r[5]} | {r[6]} | [`{r[7]}`]({r[7]}/spec.md) |" for r in index]
    (out / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"{REPO}: {len(index)} specs en el índice, {created} creadas o regeneradas")
    return 0


if __name__ == "__main__":
    sys.exit(main())
